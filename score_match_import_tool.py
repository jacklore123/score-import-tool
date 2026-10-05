# -*- coding: utf-8 -*-
"""
成绩智能匹配导入工具（批量版）
==============================

选择一份或多份“原始成绩表”，再选择一份或多份“系统表格”，自动配对并逐份导出。
单份文件也在这里处理：各选一份即可。

匹配逻辑：按“姓名”列精确匹配（等价 Excel VLOOKUP(..., 0)）。
“缺考”或未匹配到：成绩留空，并在“备注”列写“缺考”。
"""

import base64
import datetime
import hashlib
import os
import queue
import re
import shutil
import sys
import threading
import traceback
import tempfile
import zipfile
from difflib import SequenceMatcher

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    _HAS_DND = True
except ImportError:
    _HAS_DND = False


_TEMP_DIRS = []
_ZIP_FILENAME_WARNINGS = []


def _register_temp_dir(path):
    if path not in _TEMP_DIRS:
        _TEMP_DIRS.append(path)


def _cleanup_temp_dirs():
    while _TEMP_DIRS:
        d = _TEMP_DIRS.pop()
        shutil.rmtree(d, ignore_errors=True)


APP_ICON_B64 = "iVBORw0KGgoAAAANSUhEUgAAAQAAAAEACAYAAABccqhmAAAK4UlEQVR4nO3dP4sk1xXG4V6hSJG+gJwLgWKhQKBAAoNjBwpUSqdTbbihQintTVUOHDheWLCCBQdGsUEot76Ao40MNoVoqbemurv+3Fv3Pef9PWCMx7Mz3bfOeevc6q6ewwEAAPh50voBZHT69pv/Xfv/jk+fseasnwyKcafGHyMIWD8FBMDOjT9GELB+LREAFZr/T3/87Or3v3j5w6OvOYcA69eWbeHVKN5bjX8vCBxDgPVr763WD8CxeKe+f8s2IiLWT4PdWad08S5t/HvTgMMkwPrpYAJoXLzjn5N9EmD9tBAAjYvXKQRYPz0EwEx7N2W2EGD9NBEAK5Q++9f+uWpYPx0EQOPR1WErwPrpIgCEijdjCLB+2ggAwBgBIHb2yjQFsH76CADAGAEgePbKMAWwfjEQAIAxAkD07BV5CmD94iAAAGPp7zxrdfZ679//ffS1X/7wdvo7Blm/WJgAKphq/ltfB+vXCgFQ2L0mJwRYPyUEQMHxdW5zrwmBCBcDWb94CIBCljY1kwDrp4AAKGBtMxMCrF9rBIDga9cRtwGsX0wEwEZbz+LuUwDr1xYBINC8riHA+rVHAAQZX5W3AaxfXAQAYIwAEBnb3bYBrJ8G+wCINL4qbgNYv9jsA0DpbO0yBbB+OggAwBgBABgjABbu/2uP6XN/vtr1CtYvJusAaH0BLfrzYP3isw4AwB0BABgjAARfpsv6ciDrp4cAELygFu1xqzwOl8ddkm0AZLmA1er5sH452AYAAAIAsMYEIHphLtuFQNZPEwEAGCMAAGOWARDxHnalzwZg/fKwDAAAvyIAAGMEgPAV+SyvBLB+uggAwBgBABgjAABjBABgjAAAjNkFQKY3sbR4MxDrl4tdAAD4HQEAGCMAxN+M0/r3R3/8rX+/OgIAMEYAAMYIAMAYAQAYIwAAYwQAYIwAAIy93foBoI4abwk+Pn32pPTPRFsEQGAvXv4w+fVa9wOMf+7l789yX4UbAiBwo0d4jA8ffLr7Y8F8BICg5z+9OmR9LgSCFgIgadN3X315KK3//i9Fnydh0B4BELjpazT52t+3JhwIg/YIgCBNv3ezl3h8S0KBMGiDABBtevWGX/Mc5gYCYbAfAkCk8TM0fI1AOK8h1wvqIAAaNr5D0899/vfCgCCogwAohKbfJwzYHpRFAGxE49cLA6aC+giAis3vPuLvORVwjWAdAmAFGl9vKuAawToEwAI0fnsEQVkEwAw0vh6CoAzL+7uX/HmrW3fiscfXcesawZJjfDT7zAM+EehGUdD8cdwK43vH0hlbgAk0fs5twXBc+eCSNxEAM5ufcT9HEBACb2ILMGNMpPljunbc2BL8jgmAs35qTAO32U8AnPU93JoGnNkHwBRG/pw4ro+xBbhAgeQ390YjF3YTwLXPzKf5vVw73qdKf1NBlVUA0Py41BECPgFA82NKZx4CFgFA8+OWzjgE0gcAzY85OtMQSB0AND+W6AxDIHUATOFqP6gPgwCYSm2aH3N0E5NA1ikgZQDQ/NiqMwmBdAFA86OUziAEUgUAzY/SuuQhkCYAaH7U0iUOgTQBMMYFP1BPJgGQJY0RyylB3YUPAEZ/7KVLuBUIHwBjjP6gvkwCIHr6IodT4DoMGwCM/milS7QVCBsAY4z+oN5MAiBq2iK3U8C6DBkAY5z9Qd2ZBEDElIWPU7D6DBcAY5z9Qf2ZBMA4XWl+KOhGrwpEmgJCBQAA0wDg7A9lXdApgD8N1sBHn3z+6Gs//uPvBzesQ3tPDgFkeNffVLFfkzkMMq9DP/H3Bo9Pn0n3WJgtgEvzr/n+KLKvQxesLkMEQJS9VOkijlb897iuw0m8fuUDIHLKbi3e6MXvuA5doPoMGQBRLC3af/3082//2fJz1LAO2qQDQH18KmXc9OP/7SLrOpyE61g6AKKOV0vOeteK/PLrUacA13XogtRpuAAAYBIAymMTkKWeZQMg6lhVa0yNMv6eua9DF6RewwQAgPIIgMY+/OD9RV/PinVoQzIAVPdLexW/W/O7rMNJsK5D3A0YZT+1RbZiXyvTOnRffTl5g5ASyQkAwD4IAMAYAVBYrXvYo90bzzrEQAAAxuQCQPFKKZC1vuUCIMMrAKXH32jj/xnrcJCvX/kAAFAPASB+9ot69j9jHbQRAMLFH735z1gHXQSAaPFnaf4z1kETASBY/Nma/4x10BPiXoBsxe/8F3FYBy0EQAMuzX4P69AeWwDAGAFg7p2//bP1Q0BDabcAz0+n1g/h8HA8HiI0//Dfr//88UEZx7MOJgDAGAFgajz6sxXwlHYLoD5+YxmOZx1MAIaune2ZAvzIB4D6hypGc6/JCQGv+pULgOPTZ09aPwbApb7lAgD1zD27MwX4IABMLG1qQsBD2lcBSnvn5/88v/c9r99/92GfRwOORxkEwIYCu/Vvvnv11weVl67Wns2V3iE4vBPw60+/eL7lGBLQQQNguJK614crnotmaOBSP4vC276Gw1sBth6Trw9f7Ho8evFXAGQDYLhSuvXjkze8d/yhRhG3CoGte/mWU8CaKeyW3wLk1bp//7BxolN7BUA2AFQsHTkVQ0BlhN/a/FuPRYmJLqPUATA3sUufaW79HrYD89ap9NpfC5DXM0NZ4W5E65cBl+6n1Jq/1e+LRvV4PCwc/yPs/6UDQHG/BGSrZ9kAAFBfqACIMlaprdnaddvyb531gdZMOgBUxya3QoxU0IqOwnUsHQB72PuqPK8CaK3Pa/O3b4cLgBpno72KYO9iu3z35NJ1u/z+vf/EdeTj0QebluQDYK/xqXbRKZxp5hanQhFnOR5H4fE/RADsWaC1iqJl8289e+999o98PHqB4FxKOp0uje8NqF2YJd6QonDWv1acU+s353taiXA8+tH6qZ/9078VuGSxRL//fGjmywK9d7ZSav6Mx0OFfEK1nAIymjOmsq4eZ/8BE4CZc3NPBQGN7ydESl1iCoCaPujZP+yrANGvvCKPPnj9hQuASOkKP8dg9RkuADKmMGLqE9RdyACIlrLwcAxYlyEDIGsaI44+Sb29lSltsxwUaOsn6izi2T90AERedORyDFyHoQNgClMAqC+jAGArgL30iUb/NAGQ4SAgpmOCuksRAFPYCoB6MgoAtgKopU84+qcLgAEhgNL6xM2fLgAGhABK6ZM3f8oAGBAC2Ko3aP60ATAgBLBWb9L8qQPgGl4dAPVhEgDXUpsQwJK6OCY9+6cPgAEhgDl6w+a3CIABIYBbetPmtwmAASGAKb1x81sFwIAQwKXevPkHNk/03seLD/hcfB9OL/XdYveE74XAgCDIi7O+8RZgjC2BF5r/MesAGBACHmj+adZbgDnbgQFbgrhuvenraLjnH7NfgDGuC+TBWf8+AmAC00BsnPXnIwBWhMCAbYGee/d4MPI/RgDcwTQQA2f9dQiAGZgGdHHW34YAWIAg0EHjl0EArEAQtEPjl0UAbEAQ7IfGr4MAqBwCZ7xqsNzcT27i6v56BEAhBEE5NP5+CIAKCIPlaPo2CACBIHDdIiz5cFbG/DoIALEgyBwIaz6NmcaviwAIEAZRA2Htx6/T9PshAAKGgWIobP1bCzR9GwRAkjDYIxxK/0EVmr49AsAgDJTQ9FoIgAAiBwINr40ACEwpGGj0mAiApGqEA00OAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAIdf/R+DSQABXGv9vQAAAABJRU5ErkJggg=="


def _openpyxl():
    import openpyxl
    return openpyxl


def _xlrd():
    import xlrd
    return xlrd


def _preload_dependencies():
    _openpyxl()
    _xlrd()


# ---------------------------------------------------------------------------
# 通用工具函数
# ---------------------------------------------------------------------------

def get_application_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def strip_name(value):
    if value is None:
        return ""
    return str(value).strip()


def is_absent_marker(value):
    """判断成绩是否为“缺考/空/NaN/斜杠”等应留空并在备注标记缺考的情况。"""
    if value is None:
        return True
    if isinstance(value, float) and value != value:
        return True
    if isinstance(value, str):
        return value.strip() in ("", "缺考", "/", "／", "\\", "-", "—")
    return False


def normalize_score(value):
    """数字字符串转纯数值，其他文本保持原样；缺考等标记返回 None。"""
    if is_absent_marker(value):
        return None
    if isinstance(value, str):
        text = value.strip()
        try:
            return float(text)
        except ValueError:
            return text
    return value


def read_first_sheet(path):
    if not os.path.isfile(path):
        raise FileNotFoundError("文件不存在：{}".format(path))

    ext = os.path.splitext(path)[1].lower()
    if ext not in (".xlsx", ".xls"):
        raise ValueError("不支持的文件格式：{}，请选择 .xlsx 或 .xls 文件".format(ext))

    try:
        if ext == ".xlsx":
            openpyxl = _openpyxl()
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.worksheets[0]
                return [list(row) for row in ws.iter_rows(values_only=True)]
            finally:
                wb.close()
        else:
            xlrd = _xlrd()
            wb = xlrd.open_workbook(path)
            ws = wb.sheet_by_index(0)
            return [
                [ws.cell_value(r, c) for c in range(ws.ncols)]
                for r in range(ws.nrows)
            ]
    except Exception as exc:
        raise ValueError(
            "读取 Excel 失败：{}。\n请确认文件没有损坏、没有加密，且扩展名与实际格式一致。".format(exc)
        )


def extract_class_label(text):
    """只从最后一级目录名或文件名里提取班级号，如“高一9”返回“9班”。"""
    if not text:
        return None
    base = os.path.basename(os.path.normpath(text))
    if os.path.isdir(text):
        name = base
    else:
        name = os.path.splitext(base)[0]
    m = re.match(r"^高一[（(]?(\d{1,3})[）)]?班?$", name)
    if not m:
        m = re.match(r"^(\d{1,3})班$", name)
    if m:
        return "{}班".format(int(m.group(1)))
    return None


def extract_class_from_table_name(path):
    """从表文件名里提取班级号，如“原始成绩_1班”返回“1班”；匹配不到返回 None。"""
    name = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r"高一[（(]?(\d{1,3})[）)]?", name)
    if m:
        return "{}班".format(int(m.group(1)))
    m = re.search(r"(\d{1,3})\s*班", name)
    if m:
        return "{}班".format(int(m.group(1)))
    return None


def make_label(xlsx_display, container_item):
    """按顺序取第一个非 None 值作为 label。"""
    return (
        extract_class_from_table_name(xlsx_display)
        or extract_class_from_table_name(container_item)
        or extract_class_label(container_item)
        or os.path.splitext(os.path.basename(xlsx_display))[0]
    )


def parse_drop_paths(event_data):
    """把 tkinterdnd2 的拖拽字符串解析成路径列表。"""
    items = re.findall(r"\{[^}]*\}|[^\s]+", event_data)
    out = []
    for i in items:
        if i.startswith("{") and i.endswith("}"):
            out.append(i[1:-1])
        else:
            out.append(i)
    return out


def classify_dropped_paths(paths):
    """把拖拽进来的路径分成 (files, folders)。"""
    files = []
    folders = []
    for p in paths:
        if os.path.isdir(p):
            folders.append(p)
        elif os.path.splitext(p)[1].lower() in (".xlsx", ".xls", ".zip"):
            files.append(p)
    return files, folders


def _list_zip_files(zip_path):
    with zipfile.ZipFile(zip_path) as zf:
        out = []
        for i in zf.infolist():
            if i.is_dir():
                continue
            orig = i.filename
            display = orig
            if not (i.flag_bits & 0x800):
                try:
                    display = orig.encode("cp437").decode("gbk")
                except Exception:
                    _ZIP_FILENAME_WARNINGS.append(
                        "压缩包 {} 中存在编码异常的文件名：{}".format(os.path.basename(zip_path), orig)
                    )
            if display != orig:
                _ZIP_FILENAME_WARNINGS.append(
                    "压缩包 {} 中文件名 {} 已解码为 {}".format(os.path.basename(zip_path), orig, display)
                )
            out.append((orig, display))
        return out


def _has_name_and_score(path, name_column, score_column):
    try:
        rows = read_first_sheet(path)
    except Exception:
        return False
    if not rows:
        return False
    header_idx = detect_header_row(rows, name_column)
    header = [strip_name(c) for c in rows[header_idx]]
    if find_column(header, name_column) is None:
        return False
    sub_header = None
    if header_idx + 1 < len(rows):
        sub_header = [strip_name(c) for c in rows[header_idx + 1]]
    return find_score_index(header, sub_header, score_column) is not None


def find_score_file_by_scan(candidates, name_column, score_column):
    """遍历候选 Excel 文件，找出真正同时包含姓名列和成绩列的那一个。"""
    hits = [c for c in candidates if _has_name_and_score(c, name_column, score_column)]
    if len(hits) == 1:
        return hits[0]
    if len(hits) > 1:
        raise ValueError(
            "发现多个同时包含姓名和成绩列的表格，请手动指定或只保留一个：{}".format(
                "、".join(os.path.basename(c) for c in hits)
            )
        )
    return None


def collect_excel_from_zip(zip_path):
    """解压 zip，返回所有 .xlsx/.xls 的 [(解压后路径, display名)]。"""
    files = _list_zip_files(zip_path)
    if not files:
        raise ValueError("压缩包为空：{}".format(os.path.basename(zip_path)))
    tmp_dir = tempfile.mkdtemp(prefix="orig_zip_")
    _register_temp_dir(tmp_dir)
    excel = [(o, d) for o, d in files if os.path.splitext(d)[1].lower() in (".xlsx", ".xls")]
    with zipfile.ZipFile(zip_path) as zf:
        for o, d in excel:
            zf.extract(o, tmp_dir)
    return [(os.path.join(tmp_dir, o), d) for o, d in excel]


def expand_original_item(item, name_column="姓名", score_column="成绩"):
    """把原始项展开成若干 (path, label)。单个文件返回一项；zip/文件夹可能返回多项。"""
    if not os.path.isdir(item):
        ext = os.path.splitext(item)[1].lower()
        if ext not in (".zip", ".xlsx", ".xls"):
            raise ValueError("不支持的原始文件类型：{}".format(os.path.basename(item)))
        if ext in (".xlsx", ".xls"):
            return [(item, make_label(os.path.basename(item), item))]

    pairs = []
    if os.path.isdir(item):
        for root, dirs, files in os.walk(item):
            dirs.sort()
            files.sort()
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                full = os.path.join(root, f)
                if ext in (".xlsx", ".xls"):
                    pairs.append((full, f, item))
                elif ext == ".zip":
                    try:
                        for p, d in collect_excel_from_zip(full):
                            pairs.append((p, d, full))
                    except Exception:
                        continue
    else:
        for p, d in collect_excel_from_zip(item):
            pairs.append((p, d, item))

    small = [(p, d, c) for p, d, c in pairs if "小题分" in os.path.basename(d)]
    if small:
        valid = small
    else:
        valid = [(p, d, c) for p, d, c in pairs if _has_name_and_score(p, name_column, score_column)]
        if len(valid) > 1:
            summary_kw = ("汇总", "总表", "总成绩", "总分表")
            filtered = [(p, d, c) for p, d, c in valid if not any(kw in os.path.basename(d) for kw in summary_kw)]
            if filtered:
                valid = filtered
        analysis_filtered = [(p, d, c) for p, d, c in valid if not any(kw in os.path.basename(d) for kw in ANALYSIS_KEYWORDS)]
        if analysis_filtered:
            valid = analysis_filtered

    if not valid:
        raise ValueError("原始项 {} 未找到含姓名和成绩列的表格".format(item))

    if len(valid) == 1:
        p, d, c = valid[0]
        return [(p, make_label(d, c))]

    return [(p, make_label(d, c)) for p, d, c in valid]


def expand_system_item(item):
    """展开系统表项：zip/文件夹可能含多个 .xlsx/.xls，返回 [(path, label)]。"""
    if not os.path.isdir(item):
        ext = os.path.splitext(item)[1].lower()
        if ext in (".xlsx", ".xls"):
            return [(item, make_label(os.path.basename(item), item))]
        if ext != ".zip":
            raise ValueError("不支持的系统表类型：{}".format(os.path.basename(item)))
        candidates = [(p, make_label(d, item)) for p, d in collect_excel_from_zip(item)]
    else:
        candidates = []
        for root, dirs, files in os.walk(item):
            dirs.sort()
            files.sort()
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                full = os.path.join(root, f)
                if ext in (".xlsx", ".xls"):
                    candidates.append((full, make_label(f, item)))
                elif ext == ".zip":
                    try:
                        for p, d in collect_excel_from_zip(full):
                            candidates.append((p, make_label(d, full)))
                    except Exception:
                        continue

    if not candidates:
        raise ValueError("系统项 {} 未找到可用的系统模板".format(item))

    first = [(p, lb) for p, lb in candidates if any(kw in os.path.basename(p) for kw in SYSTEM_TEMPLATE_KEYWORDS)]
    if first:
        return first
    second = [(p, lb) for p, lb in candidates if not any(kw in os.path.basename(p) for kw in ANALYSIS_KEYWORDS)]
    if second:
        return second
    third = [(p, lb) for p, lb in candidates if read_names(p, "姓名")]
    if third:
        return third
    raise ValueError("系统项 {} 未找到可用的系统模板".format(item))


def find_column(header, target, allow_contains=True):
    """在表头中查找目标列：先精确匹配，再开头匹配，最后包含匹配。"""
    if not target:
        return None
    cells = [strip_name(c) for c in header]
    if target in cells:
        return cells.index(target)
    for i, h in enumerate(cells):
        if h.startswith(target):
            return i
    if allow_contains:
        for i, h in enumerate(cells):
            if target in h:
                return i
    return None


SCORE_KEYWORDS = ("成绩", "分数", "总分", "得分")
SYSTEM_TEMPLATE_KEYWORDS = ("成绩导入", "导入模板", "模板", "成绩录入")
ANALYSIS_KEYWORDS = ("作答详情", "分数段", "分布", "综合指标", "试题分析", "试题", "分析", "报告", "统计", "详情", "明细")


def find_score_index(header, sub_header, score_column):
    """在表头（必要时连同子表头）中定位成绩列，兼容“物理/分数”这类两级表头。"""
    idx = find_column(header, score_column)
    if idx is not None:
        return idx
    for kw in (score_column,) + SCORE_KEYWORDS:
        idx = find_column(header, kw)
        if idx is not None:
            return idx
        if sub_header is not None:
            idx = find_column(sub_header, kw)
            if idx is not None:
                return idx
    return None


def detect_header_row(rows, name_column, max_scan=10):
    """在表格前若干行里找到真正包含姓名列的表头行。"""
    for i, row in enumerate(rows[:max_scan]):
        if row and find_column([strip_name(c) for c in row], name_column, allow_contains=False) is not None:
            return i
    return 0


def build_score_map(rows, name_column, score_column, path_label):
    """从原始成绩表中读取 姓名->成绩 映射；返回 (score_map, name_counts, header_idx)。"""
    if not rows:
        raise ValueError("{}内容为空，请检查文件。".format(path_label))
    header_idx = detect_header_row(rows, name_column)
    header = [strip_name(c) for c in rows[header_idx]]
    sub_header = None
    if header_idx + 1 < len(rows):
        sub_header = [strip_name(c) for c in rows[header_idx + 1]]
    name_idx = find_column(header, name_column)
    score_idx = find_score_index(header, sub_header, score_column)
    if name_idx is None or score_idx is None:
        missing = []
        if name_idx is None:
            missing.append(name_column)
        if score_idx is None:
            missing.append(score_column)
        raise ValueError("{}缺少列：{}。".format(path_label, "、".join(missing)))

    score_map = {}
    name_counts = {}
    for row in rows[header_idx + 1:]:
        if not row:
            continue
        name = strip_name(row[name_idx] if name_idx < len(row) else None)
        if not name:
            continue
        name_counts[name] = name_counts.get(name, 0) + 1
        if name not in score_map:
            score_map[name] = row[score_idx] if score_idx < len(row) else None
    return score_map, name_counts, header_idx


def export_standard_original(original_path, output_dir, name_column, score_column, remark_column, class_label):
    """从原始成绩表（含 zip 解压出的小题分）抽出姓名、成绩，生成标准原始表。"""
    rows = read_first_sheet(original_path)
    if not rows:
        raise ValueError("原始成绩表内容为空，请检查文件。")
    header_idx = detect_header_row(rows, name_column)
    header = [strip_name(c) for c in rows[header_idx]]
    sub_header = None
    if header_idx + 1 < len(rows):
        sub_header = [strip_name(c) for c in rows[header_idx + 1]]
    name_idx = find_column(header, name_column)
    score_idx = find_score_index(header, sub_header, score_column)
    if name_idx is None or score_idx is None:
        raise ValueError("原始成绩表缺少姓名或成绩列，无法生成标准原始表。")

    out_rows = [[name_column, score_column, remark_column]]
    for row in rows[header_idx + 1:]:
        name = strip_name(row[name_idx] if name_idx < len(row) else None)
        if not name:
            continue
        score = normalize_score(row[score_idx] if score_idx < len(row) else None)
        out_rows.append([name, score, ""])

    openpyxl = _openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in out_rows:
        ws.append(r)

    class_part = class_label if class_label else os.path.splitext(os.path.basename(original_path))[0]
    output_path = os.path.join(output_dir, "{}原始成绩表.xlsx".format(class_part))
    wb.save(output_path)
    return output_path


def _fill_xlsx(system_path, name_column, score_column, remark_column, score_map, output_path):
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(system_path)
    ws = wb.worksheets[0]
    rows_values = [[c.value for c in row] for row in ws.iter_rows()]
    if not rows_values:
        raise ValueError("系统表格内容为空，请检查文件。")

    header_idx = detect_header_row(rows_values, name_column)
    header = [strip_name(c) for c in rows_values[header_idx]]
    name_idx = find_column(header, name_column)
    if name_idx is None:
        raise ValueError("系统表格缺少“{}”列。".format(name_column))

    score_idx = find_column(header, score_column)
    remark_idx = find_column(header, remark_column)
    max_col = max(len(r) for r in rows_values) if rows_values else 0
    if score_idx is None:
        score_idx = max_col
        ws.cell(row=header_idx + 1, column=score_idx + 1, value=score_column)
        max_col += 1
    if remark_idx is None:
        remark_idx = max_col
        ws.cell(row=header_idx + 1, column=remark_idx + 1, value=remark_column)
        max_col += 1

    matched = absent = missing = overwritten = 0
    for r in range(header_idx + 1, ws.max_row):
        row_num = r + 1
        name = strip_name(ws.cell(row=row_num, column=name_idx + 1).value)
        value = score_map.get(name)
        score_cell = ws.cell(row=row_num, column=score_idx + 1)
        remark_cell = ws.cell(row=row_num, column=remark_idx + 1)
        existing = score_cell.value
        if score_cell.data_type == "f" or (existing is not None and str(existing).strip() != ""):
            overwritten += 1
        if name in score_map and is_absent_marker(value):
            score_cell.value = None
            remark_cell.value = "缺考"
            absent += 1
        elif name not in score_map:
            score_cell.value = None
            remark_cell.value = "缺考"
            missing += 1
        else:
            score_cell.value = normalize_score(value)
            matched += 1
    wb.save(output_path)
    return matched, absent, missing, overwritten, [], output_path


def _fill_xls(system_path, name_column, score_column, remark_column, score_map, output_path):
    xlrd = _xlrd()
    extra_warnings = []
    try:
        from xlutils.copy import copy as xlcopy
    except Exception:
        extra_warnings.append("xlutils 不可用，已降级为 xlsx 导出")
        return _fill_xls_fallback_xlsx(
            system_path, name_column, score_column, remark_column, score_map, output_path, extra_warnings
        )

    try:
        rb = xlrd.open_workbook(system_path, formatting_info=True)
    except (TypeError, NotImplementedError):
        rb = xlrd.open_workbook(system_path, formatting_info=False)
        extra_warnings.append("xlrd 版本不支持保留 .xls 原格式，已降级导出")

    ws = rb.sheet_by_index(0)
    rows_values = [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
    if not rows_values:
        raise ValueError("系统表格内容为空，请检查文件。")

    header_idx = detect_header_row(rows_values, name_column)
    header = [strip_name(c) for c in rows_values[header_idx]]
    name_idx = find_column(header, name_column)
    if name_idx is None:
        raise ValueError("系统表格缺少“{}”列。".format(name_column))

    score_idx = find_column(header, score_column)
    remark_idx = find_column(header, remark_column)
    score_col_added = score_idx is None
    remark_col_added = remark_idx is None
    wb = xlcopy(rb)
    out_ws = wb.get_sheet(0)
    max_col = ws.ncols
    if score_idx is None:
        score_idx = max_col
        out_ws.write(header_idx, score_idx, score_column)
        max_col += 1
    if remark_idx is None:
        remark_idx = max_col
        out_ws.write(header_idx, remark_idx, remark_column)
        max_col += 1

    matched = absent = missing = overwritten = 0
    for r in range(header_idx + 1, ws.nrows):
        name = strip_name(ws.cell_value(r, name_idx))
        value = score_map.get(name)
        existing = None if (score_col_added or score_idx >= ws.ncols) else ws.cell_value(r, score_idx)
        if existing is not None and str(existing).strip() != "":
            overwritten += 1
        if name in score_map and is_absent_marker(value):
            out_ws.write(r, score_idx, "")
            out_ws.write(r, remark_idx, "缺考")
            absent += 1
        elif name not in score_map:
            out_ws.write(r, score_idx, "")
            out_ws.write(r, remark_idx, "缺考")
            missing += 1
        else:
            out_ws.write(r, score_idx, normalize_score(value))
            matched += 1
    wb.save(output_path)
    return matched, absent, missing, overwritten, extra_warnings, output_path


def _fill_xls_fallback_xlsx(system_path, name_column, score_column, remark_column, score_map, output_path, extra_warnings):
    """xlutils 不可用时，用 xlrd 读取 .xls 值，再用 openpyxl 另存为 .xlsx。"""
    rows = read_first_sheet(system_path)
    if not rows:
        raise ValueError("系统表格内容为空，请检查文件。")
    header_idx = detect_header_row(rows, name_column)
    header = [strip_name(c) for c in rows[header_idx]]
    name_idx = find_column(header, name_column)
    if name_idx is None:
        raise ValueError("系统表格缺少“{}”列。".format(name_column))

    for col in (score_column, remark_column):
        if find_column(header, col) is None:
            header.append(col)
            for row in rows:
                row.append(None)
    score_idx = find_column(header, score_column)
    remark_idx = find_column(header, remark_column)

    def set_cell(row, index, value):
        while len(row) <= index:
            row.append(None)
        row[index] = value

    matched = absent = missing = overwritten = 0
    for row in rows[header_idx + 1:]:
        name = strip_name(row[name_idx] if name_idx < len(row) else None)
        value = score_map.get(name)
        existing = row[score_idx] if score_idx < len(row) else None
        if existing is not None and str(existing).strip() != "":
            overwritten += 1
        if name in score_map and is_absent_marker(value):
            set_cell(row, score_idx, None)
            set_cell(row, remark_idx, "缺考")
            absent += 1
        elif name not in score_map:
            set_cell(row, score_idx, None)
            set_cell(row, remark_idx, "缺考")
            missing += 1
        else:
            set_cell(row, score_idx, normalize_score(value))
            matched += 1

    openpyxl = _openpyxl()
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    new_path = os.path.splitext(output_path)[0] + ".xlsx"
    wb.save(new_path)
    return matched, absent, missing, overwritten, extra_warnings, new_path


def match_and_export(original_path, system_path, name_column, score_column, remark_column, output_dir, extra_suffix="", output_stem=None):
    """单份匹配并导出。返回 (output_path, matched, absent, missing, warnings)。"""
    original_rows = read_first_sheet(original_path)
    score_map, original_name_counts, _ = build_score_map(
        original_rows, name_column, score_column, "原始成绩表"
    )

    warnings = []
    for name, count in original_name_counts.items():
        if count > 1:
            warnings.append(
                "原始表 {} 中姓名“{}”出现 {} 次，仅使用第一次出现的成绩".format(
                    os.path.basename(original_path), name, count
                )
            )

    system_rows = read_first_sheet(system_path)
    if not system_rows:
        raise ValueError("系统表格内容为空，请检查文件。")
    sys_header_idx = detect_header_row(system_rows, name_column)
    sys_header = [strip_name(c) for c in system_rows[sys_header_idx]]
    sys_name_idx = find_column(sys_header, name_column)
    if sys_name_idx is None:
        raise ValueError("系统表格缺少“{}”列，无法按姓名匹配。".format(name_column))

    system_name_counts = {}
    for row in system_rows[sys_header_idx + 1:]:
        if not row:
            continue
        name = strip_name(row[sys_name_idx] if sys_name_idx < len(row) else None)
        if name:
            system_name_counts[name] = system_name_counts.get(name, 0) + 1
    for name, count in system_name_counts.items():
        if count > 1:
            warnings.append(
                "系统表 {} 中姓名“{}”出现 {} 次，将都赋予同一成绩".format(
                    os.path.basename(system_path), name, count
                )
            )

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(system_path))[0]
    ext = os.path.splitext(system_path)[1].lower()
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:21]
    orig_short = hashlib.md5(original_path.encode("utf-8")).hexdigest()[:6]
    if output_stem:
        output_name = output_stem + ext
    else:
        suffix_part = extra_suffix or ""
        output_name = "{}_匹配结果{}_{}_{}{}".format(base_name, suffix_part, timestamp, orig_short, ext)
    output_path = os.path.join(output_dir, output_name)
    i = 1
    while os.path.exists(output_path):
        stem = os.path.splitext(output_name)[0]
        output_path = os.path.join(output_dir, "{}_({}){}".format(stem, i, ext))
        i += 1

    if ext == ".xls":
        matched, absent, missing, overwritten, extra_warnings, output_path = _fill_xls(
            system_path, name_column, score_column, remark_column, score_map, output_path
        )
    else:
        matched, absent, missing, overwritten, extra_warnings, output_path = _fill_xlsx(
            system_path, name_column, score_column, remark_column, score_map, output_path
        )

    warnings.extend(extra_warnings)
    if overwritten > 0:
        warnings.append(
            "系统表 {} 的成绩列原有 {} 个非空值，已被本次匹配覆盖".format(
                os.path.basename(system_path), overwritten
            )
        )

    return output_path, matched, absent, missing, warnings


def plan_batch_pairs(original_paths, system_paths):
    """根据数量关系自动配对，返回 [(原始表路径, 系统表路径, 文件名后缀)]。"""
    original_paths = list(original_paths)
    system_paths = list(system_paths)
    if not original_paths or not system_paths:
        return []

    if len(original_paths) == 1:
        return [(original_paths[0], s, "") for s in system_paths]
    if len(system_paths) == 1:
        return [(o, system_paths[0], "") for o in original_paths]
    if len(original_paths) == len(system_paths):
        return list(zip(original_paths, system_paths, [""] * len(original_paths)))

    pairs = []
    system_usage = {}
    for o in original_paths:
        for s in system_paths:
            used = system_usage.get(s, 0)
            suffix = "" if used == 0 else "_{}".format(used + 1)
            pairs.append((o, s, suffix))
            system_usage[s] = used + 1
    return pairs




CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}


def _parse_cn_number(text):
    """把“九”“十二”等中文数字转成阿拉伯数字。"""
    if text in CN_NUM:
        return CN_NUM[text]
    if "十" in text:
        parts = text.split("十")
        tens = CN_NUM.get(parts[0], 1) if parts[0] else 1
        ones = CN_NUM.get(parts[1], 0) if parts[1] else 0
        return tens * 10 + ones
    return None


def extract_class_tokens(filename):
    """从文件名中提取班级关键词，如“17班”“九班”。"""
    name = os.path.splitext(os.path.basename(filename))[0]
    tokens = set()
    for m in re.finditer(r"(\d{1,3})\s*班", name):
        tokens.add("{}班".format(int(m.group(1))))
    for m in re.finditer(r"([一二三四五六七八九十]{1,3})\s*班", name):
        n = _parse_cn_number(m.group(1))
        if n:
            tokens.add("{}班".format(n))
    return tokens


def _filename_key(path):
    """去掉扩展名，小写，用于比较相似度。"""
    return os.path.splitext(os.path.basename(path))[0].lower()


def filename_match_score(original_path, system_path):
    """文件配对打分：班级关键词越一致越高，其次比较文件名字符串相似度。"""
    o_tokens = extract_class_tokens(original_path)
    s_tokens = extract_class_tokens(system_path)
    score = 0.0

    if o_tokens and s_tokens:
        common = o_tokens & s_tokens
        if common:
            score += 120.0 + 30.0 * len(common)
        else:
            score -= 30.0

    similarity = SequenceMatcher(None, _filename_key(original_path), _filename_key(system_path)).ratio()
    score += similarity * 100.0
    return score


def read_names(path, name_column):
    """读取一个表格的姓名列，返回姓名集合（用于按名单自动配对）。"""
    try:
        rows = read_first_sheet(path)
        if not rows:
            return set()
        header_idx = detect_header_row(rows, name_column)
        header = [strip_name(c) for c in rows[header_idx]]
        idx = find_column(header, name_column)
        if idx is None:
            return set()
        names = set()
        for row in rows[header_idx + 1:]:
            if not row:
                continue
            name = strip_name(row[idx] if idx < len(row) else None)
            if name:
                names.add(name)
        return names
    except Exception:
        return set()


def name_overlap_score(original_names, system_names):
    """按系统表名单被原始表覆盖的比例打分：文件名完全无关也能正确配对。"""
    if not system_names:
        return 0.0
    common = len(original_names & system_names)
    return common / float(len(system_names))


def plan_batch_pairs_auto(original_paths, system_paths, name_column="姓名", min_score=0.45):
    """自动按文件内学生名单配对；名单重叠度过低的系统表不生成文件。返回 (pairs, unmatched, ambiguous)。"""
    original_paths = list(original_paths)
    system_paths = list(system_paths)
    if not original_paths or not system_paths:
        return [], [], []

    if len(original_paths) == 1:
        original_names = read_names(original_paths[0], name_column)
        pairs = []
        unmatched = []
        for s in system_paths:
            system_names = read_names(s, name_column)
            if name_overlap_score(original_names, system_names) >= min_score:
                pairs.append((original_paths[0], s, ""))
            else:
                unmatched.append(s)
        return pairs, unmatched, []

    if len(system_paths) == 1:
        system_names = read_names(system_paths[0], name_column)
        candidates = []
        for o in original_paths:
            original_names = read_names(o, name_column)
            sc = name_overlap_score(original_names, system_names)
            if sc >= min_score:
                candidates.append((sc, o))
        if not candidates:
            return [], [system_paths[0]], []
        candidates.sort(key=lambda x: x[0], reverse=True)
        best_o = candidates[0][1]
        ambiguous = [o for sc, o in candidates[1:]]
        return [(best_o, system_paths[0], "")], [], ambiguous

    orig_name_sets = [read_names(o, name_column) for o in original_paths]
    sys_name_sets = [read_names(s, name_column) for s in system_paths]

    def pair_score(oi, si):
        content = name_overlap_score(orig_name_sets[oi], sys_name_sets[si]) * 1000.0
        fname = filename_match_score(original_paths[oi], system_paths[si]) * 0.15
        return content + fname

    if len(original_paths) == len(system_paths):
        used_original = set()
        pairs = []
        unmatched = []
        for si in range(len(system_paths)):
            best_oi, best_score = None, -1.0
            for oi in range(len(original_paths)):
                if oi in used_original:
                    continue
                sc = pair_score(oi, si)
                if sc > best_score:
                    best_oi, best_score = oi, sc
            if best_oi is None:
                unmatched.append(system_paths[si])
                continue
            content = name_overlap_score(orig_name_sets[best_oi], sys_name_sets[si])
            if content < min_score:
                unmatched.append(system_paths[si])
            else:
                pairs.append((original_paths[best_oi], system_paths[si], ""))
                used_original.add(best_oi)
        return pairs, unmatched, []

    pairs = []
    unmatched = []
    used_original = set()
    for si in range(len(system_paths)):
        best_oi, best_score = None, -1.0
        for oi in range(len(original_paths)):
            if oi in used_original:
                continue
            sc = pair_score(oi, si)
            if sc > best_score:
                best_oi, best_score = oi, sc
        if best_oi is None:
            unmatched.append(system_paths[si])
            continue
        content = name_overlap_score(orig_name_sets[best_oi], sys_name_sets[si])
        if content < min_score:
            unmatched.append(system_paths[si])
        else:
            pairs.append((original_paths[best_oi], system_paths[si], ""))
            used_original.add(best_oi)
    return pairs, unmatched, []

def match_batch(original_paths, system_paths, name_column, score_column, remark_column, output_dir, log_callback=None, auto_pair=False, orig_labels=None):
    """批量匹配。逐份执行，单份失败不影响后续文件。"""
    if auto_pair:
        pairs, unmatched, ambiguous = plan_batch_pairs_auto(original_paths, system_paths, name_column)
    else:
        pairs = plan_batch_pairs(original_paths, system_paths)
        unmatched = []
        ambiguous = []
    if not pairs and not unmatched and not ambiguous:
        _cleanup_temp_dirs()
        raise ValueError("没有可匹配的文件，请先选择原始成绩表和系统表格。")
    if not pairs and (unmatched or ambiguous):
        _cleanup_temp_dirs()
        raise ValueError("没有找到可匹配的学生名单，请检查姓名列是否一致。")

    def log(msg):
        if log_callback:
            log_callback(msg)

    zip_warnings = list(_ZIP_FILENAME_WARNINGS)
    if zip_warnings:
        for w in zip_warnings:
            log("警告：" + w)
        _ZIP_FILENAME_WARNINGS.clear()

    if auto_pair and pairs:
        log("已按文件内学生名单自动配对：")
        for _o, _s, _suf in pairs:
            log("  {} ↔ {}".format(os.path.basename(_o), os.path.basename(_s)))
    if unmatched:
        log("以下系统表格未找到对应的原始成绩表，已跳过（不会生成全 0 数据）：")
        for _s in unmatched:
            log("  缺少：{}".format(os.path.basename(_s)))
    if ambiguous:
        log("系统表 {} 有多份原始表可匹配，已选择重叠度最高的一份".format(os.path.basename(system_paths[0])))
        for _o in ambiguous:
            log("  可选原始表：{}".format(os.path.basename(_o)))

    ok = 0
    fail = 0
    results = []
    standard_results = []
    all_warnings = []
    all_warnings.extend(zip_warnings)
    total = len(pairs)
    log("共计划生成 {} 份结果文件。".format(total))

    if orig_labels:
        os.makedirs(output_dir, exist_ok=True)
        for orig in list(dict.fromkeys(original_paths)):
            label = orig_labels.get(orig)
            try:
                std_path = export_standard_original(
                    orig, output_dir, name_column, score_column, remark_column, label
                )
                standard_results.append(std_path)
                log("已生成标准原始成绩表：{}".format(os.path.basename(std_path)))
            except Exception as exc:
                log("生成标准原始成绩表失败：{}（{}）".format(os.path.basename(orig), exc))

    for index, (orig, sys_file, suffix) in enumerate(pairs, start=1):
        log("正在处理 {}/{}：{} → {}".format(
            index, total, os.path.basename(orig), os.path.basename(sys_file)
        ))
        try:
            label = orig_labels.get(orig) if orig_labels else None
            output_stem = "{}成绩导入".format(label) if label else None
            output_path, matched, absent, missing, warnings = match_and_export(
                orig, sys_file, name_column, score_column, remark_column, output_dir,
                extra_suffix=suffix, output_stem=output_stem
            )
            ok += 1
            results.append(output_path)
            log("完成：{}（成功匹配 {} 条，缺考/未匹配置空并标缺考共 {} 条）".format(
                os.path.basename(output_path), matched, absent + missing
            ))
            for warning in warnings:
                log("警告：" + warning)
                all_warnings.append(warning)
        except Exception as exc:
            fail += 1
            log("失败：{}（{}）".format(os.path.basename(sys_file), exc))

    log("批量处理结束：成功 {} 份，失败 {} 份，跳过 {} 份。".format(ok, fail, len(unmatched)))
    _cleanup_temp_dirs()
    return {
        "total": total, "ok": ok, "fail": fail, "results": results,
        "standard_results": standard_results, "unmatched": unmatched,
        "ambiguous": ambiguous, "warnings": all_warnings,
    }


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

class GradeMatchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("成绩智能匹配导入工具 · 开发者 tjy")
        self.root.geometry("820x620")
        self.root.minsize(720, 540)

        self._icon_image = None
        self._set_icon()

        self.name_column = tk.StringVar(value="姓名")
        self.score_column = tk.StringVar(value="成绩")
        self.remark_column = tk.StringVar(value="备注")
        self.output_dir = tk.StringVar()

        self.message_queue = queue.Queue()
        self.running = False
        self.preloading = True

        self._build_ui()
        self.status_label.configure(text="正在加载…")
        self._log("程序已启动，正在后台加载数据组件…")
        threading.Thread(target=self._preload_worker, daemon=True).start()
        self.root.after(100, self._poll_queue)

    def _set_icon(self):
        try:
            self._icon_image = tk.PhotoImage(data=APP_ICON_B64)
            self.root.iconphoto(True, self._icon_image)
        except tk.TclError:
            self._icon_image = None

    def _preload_worker(self):
        try:
            _preload_dependencies()
            self.message_queue.put(("preload_done",))
        except Exception as exc:
            self.message_queue.put(("preload_error", str(exc)))

    def _build_ui(self):
        padding = {"padx": 10, "pady": 6}
        main = ttk.Frame(self.root)
        main.pack(fill="both", expand=True, **padding)

        pick_frame = ttk.Frame(main)
        pick_frame.pack(fill="x", pady=(0, 10))
        pick_frame.columnconfigure(0, weight=1)
        pick_frame.columnconfigure(1, weight=1)

        left = ttk.LabelFrame(pick_frame, text="原始成绩表（可多选）")
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        right = ttk.LabelFrame(pick_frame, text="系统表格（可多选）")
        right.grid(row=0, column=1, sticky="nsew", padx=(6, 0))

        ttk.Label(left, text="支持拖拽 Excel / zip / 文件夹到下方列表框", foreground="#888", font=("Segoe UI", 8)).pack(
            anchor="w", padx=8, pady=(4, 0)
        )
        ttk.Label(right, text="支持拖拽 Excel / zip / 文件夹到下方列表框", foreground="#888", font=("Segoe UI", 8)).pack(
            anchor="w", padx=8, pady=(4, 0)
        )

        self.original_list = tk.Listbox(left, selectmode="extended", height=8)
        self.original_list.pack(fill="both", expand=True, padx=8, pady=(8, 4))
        self.system_list = tk.Listbox(right, selectmode="extended", height=8)
        self.system_list.pack(fill="both", expand=True, padx=8, pady=(8, 4))

        ttk.Button(left, text="添加原始表/压缩包", command=self.add_original_files).pack(
            fill="x", padx=8, pady=3
        )
        ttk.Button(left, text="添加原始文件夹", command=self.add_original_folder).pack(
            fill="x", padx=8, pady=3
        )
        ttk.Button(left, text="移除选中/清空", command=lambda: self.remove_selected(self.original_list)).pack(
            fill="x", padx=8, pady=3
        )
        ttk.Button(right, text="添加系统表/压缩包", command=self.add_system_files).pack(
            fill="x", padx=8, pady=3
        )
        ttk.Button(right, text="添加系统文件夹", command=self.add_system_folder).pack(
            fill="x", padx=8, pady=3
        )
        ttk.Button(right, text="移除选中/清空", command=lambda: self.remove_selected(self.system_list)).pack(
            fill="x", padx=8, pady=3
        )

        self.auto_pair = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            main,
            text="自动按文件内学生名单配对（文件名不需要一致）",
            variable=self.auto_pair,
        ).pack(anchor="w", pady=(0, 4))

        hint = ("关闭“自动配对”后：1 份原始表匹配多份系统表；多份原始表匹配 1 份系统表；"
                "数量相同按顺序配对；数量都大于 1 且不同则两两组合。")
        ttk.Label(main, text=hint, foreground="#666").pack(anchor="w", pady=(0, 8))

        column_frame = ttk.LabelFrame(main, text="列名配置（默认姓名 / 成绩 / 备注，可按实际表头修改）")
        column_frame.pack(fill="x", pady=(0, 10))
        ttk.Label(column_frame, text="姓名列名：").grid(row=0, column=0, sticky="w", padx=10, pady=8)
        ttk.Entry(column_frame, textvariable=self.name_column, width=18).grid(
            row=0, column=1, sticky="w", padx=4, pady=8
        )
        ttk.Label(column_frame, text="成绩列名：").grid(row=0, column=2, sticky="w", padx=10, pady=8)
        ttk.Entry(column_frame, textvariable=self.score_column, width=18).grid(
            row=0, column=3, sticky="w", padx=4, pady=8
        )
        ttk.Label(column_frame, text="备注列名：").grid(row=1, column=0, sticky="w", padx=10, pady=8)
        ttk.Entry(column_frame, textvariable=self.remark_column, width=18).grid(
            row=1, column=1, sticky="w", padx=4, pady=8
        )

        output_frame = ttk.Frame(main)
        output_frame.pack(fill="x", pady=(0, 10))
        ttk.Button(output_frame, text="选择导出文件夹", command=self.choose_output_dir).pack(side="left")
        ttk.Entry(output_frame, textvariable=self.output_dir, state="readonly").pack(
            side="left", fill="x", expand=True, padx=10
        )

        action_frame = ttk.Frame(main)
        action_frame.pack(fill="x", pady=(0, 10))
        self.start_button = ttk.Button(action_frame, text="开始匹配", command=self.start_match)
        self.start_button.pack(side="left", padx=(0, 10))
        self.status_label = ttk.Label(action_frame, text="就绪")
        self.status_label.pack(side="left", padx=10)

        log_frame = ttk.LabelFrame(main, text="运行状态")
        log_frame.pack(fill="both", expand=True)
        self.log_text = scrolledtext.ScrolledText(
            log_frame, height=12, state="disabled", wrap="word",
            font=("Consolas", 9), background="#0f172a", foreground="#dbeafe"
        )
        self.log_text.pack(fill="both", expand=True, padx=10, pady=10)

        ttk.Label(main, text="开发者：tjy", foreground="#888").pack(
            side="bottom", anchor="e", pady=(4, 2)
        )
        self._setup_dnd()

    def choose_output_dir(self):
        path = filedialog.askdirectory(title="选择导出文件夹")
        if path:
            self.output_dir.set(path)
            self._log("导出文件夹：{}".format(path))

    def add_original_files(self):
        paths = filedialog.askopenfilenames(
            title="选择多个原始成绩表（可含 zip 压缩包）",
            filetypes=[("Excel 或压缩包", "*.xlsx *.xls *.zip"), ("所有文件", "*.*")],
        )
        for p in paths:
            if p not in self.original_list.get(0, "end"):
                self.original_list.insert("end", p)
        self._log("已添加 {} 份原始成绩表/压缩包。".format(len(paths)))

    def add_original_folder(self):
        folder = filedialog.askdirectory(title="选择原始成绩文件夹")
        if folder and folder not in self.original_list.get(0, "end"):
            self.original_list.insert("end", folder)
            self._log("已添加原始文件夹：{}".format(folder))

    def add_system_files(self):
        paths = filedialog.askopenfilenames(
            title="选择多个系统表格（可含 zip 压缩包）",
            filetypes=[("Excel 或压缩包", "*.xlsx *.xls *.zip"), ("所有文件", "*.*")],
        )
        for p in paths:
            if p not in self.system_list.get(0, "end"):
                self.system_list.insert("end", p)
        self._log("已添加 {} 份系统表格/压缩包。".format(len(paths)))

    def add_system_folder(self):
        folder = filedialog.askdirectory(title="选择系统文件夹")
        if folder and folder not in self.system_list.get(0, "end"):
            self.system_list.insert("end", folder)
            self._log("已添加系统文件夹：{}".format(folder))

    def remove_selected(self, listbox):
        selected = list(listbox.curselection())
        if selected:
            for index in reversed(selected):
                listbox.delete(index)
        else:
            listbox.delete(0, "end")

    def _setup_dnd(self):
        if not _HAS_DND:
            return
        self.original_list.drop_target_register(DND_FILES)
        self.original_list.dnd_bind("<<Drop>>", self._on_drop_original)
        self.system_list.drop_target_register(DND_FILES)
        self.system_list.dnd_bind("<<Drop>>", self._on_drop_system)

    def _on_drop_original(self, event):
        files, folders = classify_dropped_paths(parse_drop_paths(event.data))
        added = 0
        for p in files + folders:
            if p not in self.original_list.get(0, "end"):
                self.original_list.insert("end", p)
                added += 1
        self._log("拖拽添加 {} 项到原始成绩表".format(added))

    def _on_drop_system(self, event):
        files, folders = classify_dropped_paths(parse_drop_paths(event.data))
        added = 0
        for p in files + folders:
            if p not in self.system_list.get(0, "end"):
                self.system_list.insert("end", p)
                added += 1
        self._log("拖拽添加 {} 项到系统表格".format(added))

    def _log(self, message):
        timestamp = datetime.datetime.now().strftime("%H:%M:%S")
        line = "[{}] {}\n".format(timestamp, message)
        self.log_text.configure(state="normal")
        self.log_text.insert("end", line)
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def start_match(self):
        name_column = self.name_column.get().strip()
        score_column = self.score_column.get().strip()
        remark_column = self.remark_column.get().strip() or "备注"
        output_dir = self.output_dir.get().strip()

        if not name_column or not score_column:
            messagebox.showwarning("提示", "请填写姓名列名和成绩列名。")
            return
        if not output_dir:
            output_dir = os.path.join(get_application_dir(), "输出结果")

        if self.running:
            messagebox.showinfo("提示", "任务正在运行中，请稍候。")
            return

        original_items = list(self.original_list.get(0, "end"))
        system_items = list(self.system_list.get(0, "end"))
        if not original_items or not system_items:
            messagebox.showwarning("提示", "请先添加原始成绩表和系统表格。")
            return

        self.running = True
        self.start_button.configure(state="disabled")
        self.status_label.configure(text="正在匹配…")
        auto_pair = bool(self.auto_pair.get())
        self._log("开始执行批量匹配任务（{}）。".format("自动按学生名单配对" if auto_pair else "按数量规则配对"))

        worker = threading.Thread(
            target=self._worker,
            args=(original_items, system_items, name_column, score_column, remark_column, output_dir, auto_pair),
            daemon=True,
        )
        worker.start()
        self.root.after(100, self._poll_queue)

    def _worker(self, original_items, system_items, name_column, score_column, remark_column, output_dir, auto_pair):
        try:
            _ZIP_FILENAME_WARNINGS.clear()
            self.message_queue.put(("log", "正在解析原始文件/压缩包…"))
            original_paths = []
            orig_labels = {}
            for item in original_items:
                expanded = expand_original_item(item, name_column, score_column)
                if len(expanded) > 1:
                    self.message_queue.put(("log", "原始项 {} 展开为 {} 个原始表".format(item, len(expanded))))
                    for p, lb in expanded:
                        self.message_queue.put(("log", "  {}".format(os.path.basename(p))))
                for p, lb in expanded:
                    original_paths.append(p)
                    orig_labels[p] = lb
                    if len(expanded) == 1:
                        self.message_queue.put(("log", "原始项 {} 解析为 {}".format(item, p)))

            self.message_queue.put(("log", "正在解析系统表格…"))
            system_paths = []
            for item in system_items:
                try:
                    expanded = expand_system_item(item)
                    if len(expanded) > 1:
                        self.message_queue.put(("log", "系统项 {} 展开为 {} 个系统表".format(item, len(expanded))))
                        for p, d in expanded:
                            self.message_queue.put(("log", "  {}".format(os.path.basename(p))))
                    for p, d in expanded:
                        system_paths.append(p)
                except Exception as exc:
                    self.message_queue.put(("log", "系统项 {} 解析失败：{}".format(item, exc)))

            self.message_queue.put(("log", "文件解析完成，开始匹配…"))
            summary = match_batch(
                original_paths, system_paths, name_column, score_column, remark_column, output_dir,
                log_callback=lambda m: self.message_queue.put(("log", m)),
                auto_pair=auto_pair,
                orig_labels=orig_labels,
            )
            self.message_queue.put(("batch_done", summary))
        except Exception as exc:
            detail = traceback.format_exc()
            self.message_queue.put(("error", str(exc), detail))
        finally:
            _cleanup_temp_dirs()

    def _poll_queue(self):
        try:
            while True:
                message = self.message_queue.get_nowait()
                kind = message[0]

                if kind == "preload_done":
                    self.preloading = False
                    self.status_label.configure(text="就绪")
                    self._log("数据组件加载完成。")
                    continue

                if kind == "preload_error":
                    self.preloading = False
                    self._log("数据组件加载失败：{}".format(message[1]))
                    self.status_label.configure(text="组件加载失败")
                    continue

                if kind == "log":
                    self._log(message[1])

                elif kind == "batch_done":
                    summary = message[1]
                    self.status_label.configure(text="批量导出完成")
                    skipped = summary.get("unmatched", []) or []
                    if skipped:
                        shown = [os.path.basename(x) for x in skipped]
                        detail = "以下系统表格没有找到对应的原始成绩表，已跳过：\n\n"
                        if len(shown) > 5:
                            detail += "\n".join(shown[:5])
                            detail += "\n… 共 {} 个系统表未匹配".format(len(shown))
                        else:
                            detail += "\n".join(shown)
                        detail += "\n\n请补充缺少的原始成绩表后重新生成。"
                        messagebox.showwarning("有系统表格未匹配", detail)

                    ambiguous = summary.get("ambiguous", []) or []
                    if ambiguous:
                        detail = "有系统表存在多份原始表可匹配，已选择重叠度最高的一份。\n\n多出的原始表：\n"
                        detail += "\n".join(os.path.basename(x) for x in ambiguous)
                        messagebox.showwarning("多份原始表可匹配", detail)

                    cancelled = False
                    warning_list = summary.get("warnings", []) or []
                    if warning_list:
                        shown = warning_list[:5]
                        warning_text = "发现 {} 条警告，是否仍要保留已导出的结果？\n\n".format(len(warning_list))
                        warning_text += "\n".join(shown)
                        if len(warning_list) > 5:
                            warning_text += "\n... 共 {} 条警告".format(len(warning_list))
                        keep = messagebox.askyesno("导出前确认", warning_text)
                        if keep:
                            self._log("用户确认保留。")
                        else:
                            for p in summary.get("standard_results", []):
                                try:
                                    os.remove(p)
                                except Exception:
                                    pass
                            for p in summary.get("results", []):
                                try:
                                    os.remove(p)
                                except Exception:
                                    pass
                            cancelled = True
                            self._log("用户取消，已删除结果。")

                    result_desc = "已生成 {} 份（用户取消后已删除）".format(summary["ok"]) if cancelled else "成功 {} 份".format(summary["ok"])
                    messagebox.showinfo(
                        "批量完成",
                        "批量处理结束。\n{}，失败 {} 份，跳过 {} 份。".format(
                            result_desc, summary["fail"], len(skipped)
                        ),
                    )
                    self._finish()
                    return

                elif kind == "error":
                    _, error_text, detail = message
                    self._log("发生错误：{}".format(error_text))
                    self._log(detail)
                    self.status_label.configure(text="运行出错")
                    messagebox.showerror("错误", error_text)
                    self._finish()
                    return
        except queue.Empty:
            pass

        if self.running or self.preloading:
            self.root.after(100, self._poll_queue)

    def _finish(self):
        self.running = False
        self.start_button.configure(state="normal")
        self.status_label.configure(text="就绪")


def main():
    if _HAS_DND:
        root = TkinterDnD.Tk()
    else:
        root = tk.Tk()
    app = GradeMatchApp(root)
    if not _HAS_DND:
        app._log("未安装 tkinterdnd2，拖拽功能不可用，可运行：pip install tkinterdnd2")
    root.mainloop()


if __name__ == "__main__":
    main()


"""
后续手动测试建议：
- 原始表有重名学生且成绩不同：验证只取第一个成绩并产生警告。
- 原始表有重名学生且成绩相同：验证只取第一个成绩并产生警告。
- 系统表有重名学生，原始表只有一个：验证所有同名学生都得到同一成绩并产生警告。
- 原始表和系统表都有重名：验证匹配结果和警告。
- 重名学生中有缺考：验证成绩留空、备注写缺考，并产生警告。
- 重名但姓名前后有空格：验证 strip 后仍视为重名。
- 没有重名时：验证不产生警告，匹配正常。
- 重名出现在批量自动配对中：验证不影响配对，仅产生警告。
- 重名出现在单份原始表对多份系统表：验证每份系统表都得到警告。
- 重名出现在多份原始表对单份系统表：验证只使用第一个匹配的原始表并警告。
- 空姓名、None、纯空格姓名。
- 姓名列不存在、成绩列不存在。
- 成绩为文本数字、浮点数、超过 100。
- 语文 150 分制：成绩 0、50、100、105、120、135、150 均应正确保留。
- 数学 150 分制：成绩 0、50、100、105、120、135、150 均应正确保留。
- 英语 150 分制：成绩 0、50、100、105、120、135、150 均应正确保留。
- 任意科目 120 分制：成绩 0、60、100、105、115、120 均应正确保留。
- 成绩为文本 "120"：应转为数值 120。
- 成绩为 100.5：应保留浮点数。
- 成绩为 0：应保留 0，不应被当作缺考。
- 列名配置为“语文”：原始表和系统表都使用“语文”作为成绩列名，验证超过 100 分正常写入。
- 列名配置为“数学”：原始表和系统表都使用“数学”作为成绩列名，验证超过 100 分正常写入。
- 列名配置为“英语”：原始表和系统表都使用“英语”作为成绩列名，验证超过 100 分正常写入。
- 系统表已有语文/数学/英语列：验证是否覆盖配置列。
- 系统表没有配置列：验证是否追加。
- 系统表同时有语文、数学、英语等多列：只匹配配置列，其他列保持原样。
- 表头有前后空格。
- 空行、合并单元格。
- .xls 与 .xlsx 混合。
- 文件损坏、加密。
- 多 sheet，只读第一个 sheet。
- 批量配对时 1 对多、多对 1、多对多。
- min_score 阈值边界，刚好 0.45。
- 文件名乱码但名单匹配。
- 大文件性能。
- 原始表或系统表只有表头没有数据行。
- 系统表姓名列为空。
- 成绩列有公式，读取公式计算后的值。
- 输出文件名包含时间戳，不会覆盖旧文件。
- 导出目录不存在时自动创建。
- 自动配对关闭时，按数量规则配对，可能产生全 0（预期行为，仅用于对比）。
- 自动配对开启时，低于 min_score 的系统表被跳过并提示缺少对应原始表。
- 批量模式下，原始表多于系统表、系统表多于原始表、数量相同但部分不匹配。
- 系统表有多个科目列，只匹配配置列，其他列原样输出。
- 原始表有多个科目列，只读取配置列，其他列忽略。
"""
